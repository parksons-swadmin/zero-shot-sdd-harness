"""Phase-3 integration tests: SSE streaming compute + multi-sheet Excel export.

No LLM, no DB, no network — the real path is the deterministic pandas pipeline
over the committed ``ar_small.xlsx`` and the seeded ``ar_large.xlsx`` fixtures.

The headline correctness guarantee is that the STREAMED final result is
byte-identical to the non-streaming ``run_analysis`` result for the same
file+mapping — proven here at both the helper level (exact ``model_dump``
equality for small AND large) and the route level (the SSE ``result`` frame
matches ``POST /api/compute``).

``as_of`` is pinned to the fixture anchor ``2026-01-15`` — for the helper tests by
passing ``as_of`` directly, and for the route tests by the same scoped
``graph.runner.date`` monkeypatch the existing integration tests use.
"""

from __future__ import annotations

import io
import json
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api import app
from graph.runner import run_analysis, run_analysis_streaming

AS_OF = date(2026, 1, 15)
_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture
def fixed_today(monkeypatch):
    """Pin ``graph.runner.date.today()`` to the fixture anchor (buckets stay fresh)."""

    class _FixedDate(date):
        @classmethod
        def today(cls):
            return AS_OF

    monkeypatch.setattr("graph.runner.date", _FixedDate)
    return AS_OF


def _collect_progress():
    frames: list[tuple[str, int, int]] = []

    def progress(phase: str, rows_done: int, rows_total: int) -> None:
        frames.append((phase, rows_done, rows_total))

    return frames, progress


def _parse_sse(text: str) -> list[dict]:
    return [
        json.loads(line[len("data: "):])
        for line in text.splitlines()
        if line.startswith("data: ")
    ]


# --------------------------------------------------------------------------- #
# 1. Streamed result == non-streaming result — small fixture (exact equality).
# --------------------------------------------------------------------------- #
def test_streaming_equals_non_streaming_small(small_xlsx_bytes, small_mapping):
    frames, progress = _collect_progress()
    m_stream = run_analysis_streaming(
        file_bytes=small_xlsx_bytes, sheet_name=None, mapping=small_mapping,
        as_of=AS_OF, progress=progress,
    )
    m_plain = run_analysis(
        file_bytes=small_xlsx_bytes, sheet_name=None, mapping=small_mapping, as_of=AS_OF
    )
    # Byte-identical full model — totals, top customers, breakdowns, flags, all.
    assert m_stream.model_dump(mode="json") == m_plain.model_dump(mode="json")

    # Real progress: at least one frame, reaching the (small) row total.
    assert frames, "expected at least one progress frame"
    assert max(done for _, done, _ in frames) == m_stream.row_count
    assert {phase for phase, _, _ in frames} <= {"parse", "compute"}


# --------------------------------------------------------------------------- #
# 2. Streamed result == non-streaming result — large fixture (>=60,000 rows).
# --------------------------------------------------------------------------- #
def test_streaming_equals_non_streaming_large(large_fixture):
    frames, progress = _collect_progress()
    m_stream = run_analysis_streaming(
        file_bytes=large_fixture["bytes"], sheet_name=None,
        mapping=large_fixture["mapping"], as_of=AS_OF, progress=progress,
    )
    m_plain = run_analysis(
        file_bytes=large_fixture["bytes"], sheet_name=None,
        mapping=large_fixture["mapping"], as_of=AS_OF,
    )
    assert m_stream.model_dump(mode="json") == m_plain.model_dump(mode="json")

    # Progress reflects REAL rows and reaches the true total (no truncation).
    rows_total = frames[0][2]
    assert rows_total == m_stream.row_count >= 60000
    assert max(done for _, done, _ in frames) == rows_total
    # Coarse increments (~5,000 rows) — many parse frames, not one, not thousands.
    parse_frames = [f for f in frames if f[0] == "parse"]
    assert 2 <= len(parse_frames) <= 100


# --------------------------------------------------------------------------- #
# 3. SSE route: final result frame matches POST /api/compute (small).
# --------------------------------------------------------------------------- #
def test_compute_stream_route_matches_compute(client, small_xlsx_bytes, small_mapping, fixed_today):
    mapping_json = json.dumps(small_mapping.as_dict())

    stream = client.post(
        "/api/compute/stream",
        files={"file": ("ar_small.xlsx", small_xlsx_bytes, "application/octet-stream")},
        data={"mapping": mapping_json},
    )
    assert stream.status_code == 200, stream.text
    assert "text/event-stream" in stream.headers["content-type"]

    events = _parse_sse(stream.text)
    progress = [e for e in events if "phase" in e]
    results = [e for e in events if e.get("event") == "result"]
    assert progress, "expected >=1 progress frame"
    assert all(e["phase"] in ("parse", "compute") for e in progress)
    assert len(results) == 1
    streamed = results[0]["result"]

    plain = client.post(
        "/api/compute",
        files={"file": ("ar_small.xlsx", small_xlsx_bytes, "application/octet-stream")},
        data={"mapping": mapping_json},
    )
    assert plain.status_code == 200, plain.text
    expected = plain.json()["data"]

    # The streamed final result body equals the non-streaming DashboardResult.
    assert streamed == expected


def test_compute_stream_route_emits_error_frame_on_bad_mapping(client, small_xlsx_bytes):
    """A structurally invalid mapping yields a 200 stream with an error frame
    (the client then falls back), never a partial/garbage result."""
    bad = {
        "customer": "Cust Name", "invoice_no": "Invoice #", "invoice_date": "Inv. Date",
        "due_date": "No Such Column", "amount": "Balance Outstanding", "employee": "Sales Person",
    }
    resp = client.post(
        "/api/compute/stream",
        files={"file": ("ar_small.xlsx", small_xlsx_bytes, "application/octet-stream")},
        data={"mapping": json.dumps(bad)},
    )
    assert resp.status_code == 200, resp.text
    events = _parse_sse(resp.text)
    errors = [e for e in events if e.get("event") == "error"]
    assert len(errors) == 1
    assert errors[0]["code"] == "BAD_MAPPING"
    assert not any(e.get("event") == "result" for e in events)


# --------------------------------------------------------------------------- #
# 4. Excel export route: 200 + xlsx headers + read-back Summary tie-out.
# --------------------------------------------------------------------------- #
def test_export_xlsx_route_returns_workbook_that_ties_out(
    client, small_xlsx_bytes, small_mapping, expected_small, fixed_today
):
    resp = client.post(
        "/api/export/xlsx",
        files={"file": ("ar_small.xlsx", small_xlsx_bytes, "application/octet-stream")},
        data={"mapping": json.dumps(small_mapping.as_dict())},
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"] == _XLSX_MEDIA_TYPE
    cd = resp.headers["content-disposition"]
    assert "attachment" in cd
    assert "ar_small_AR_aging_2026-01-15.xlsx" in cd
    body = resp.content
    assert body and body[:2] == b"PK"

    wb = load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["Summary", "Customers", "Employees", "Flagged Rows"]

    ws = wb["Summary"]
    totals = {}
    for row in ws.iter_rows():
        if row[0].value in ("Total Outstanding", "Total Overdue"):
            totals[row[0].value] = row[1].value
    assert totals["Total Outstanding"] == expected_small["total_outstanding"]
    assert totals["Total Overdue"] == expected_small["total_overdue"]


def test_export_xlsx_route_rejects_bad_mapping(client, small_xlsx_bytes):
    """A bad mapping returns a clean 400 error, never a partial file."""
    bad = {"customer": "Cust Name"}  # missing required fields
    resp = client.post(
        "/api/export/xlsx",
        files={"file": ("ar_small.xlsx", small_xlsx_bytes, "application/octet-stream")},
        data={"mapping": json.dumps(bad)},
    )
    assert resp.status_code == 400, resp.text
    assert resp.json()["detail"]["code"] == "BAD_MAPPING"
