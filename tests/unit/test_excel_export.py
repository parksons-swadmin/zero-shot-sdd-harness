"""Unit tests for the multi-sheet Excel export (``src/tools/excel_export.py``).

The workbook is built from the SAME ``AgingMetrics`` the dashboard renders and
must carry the FULL (uncapped) data. Every assertion reads the workbook back with
openpyxl — the export cannot diverge from the API metrics.

No LLM, no DB, no network. ``as_of`` is pinned to the fixture anchor so the run
is deterministic.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from graph.runner import run_analysis
from tools.excel_export import build_workbook

AS_OF = date(2026, 1, 15)


@pytest.fixture
def small_metrics(small_xlsx_bytes, small_mapping):
    return run_analysis(
        file_bytes=small_xlsx_bytes, sheet_name=None, mapping=small_mapping, as_of=AS_OF
    )


@pytest.fixture
def small_workbook(small_metrics):
    content = build_workbook(small_metrics, source_filename="ar_small.xlsx")
    return load_workbook(io.BytesIO(content))


def _summary_value(ws, label):
    """Read the value in the cell immediately right of the row whose col-A == label."""
    for row in ws.iter_rows():
        if row[0].value == label:
            return row[1].value
    raise KeyError(label)


# --------------------------------------------------------------------------- #
# 1. The four required sheets exist by name.
# --------------------------------------------------------------------------- #
def test_workbook_has_four_named_sheets(small_workbook):
    assert small_workbook.sheetnames == ["Summary", "Customers", "Employees", "Flagged Rows"]


def test_build_workbook_returns_nonempty_bytes(small_metrics):
    content = build_workbook(small_metrics, source_filename="ar_small.xlsx")
    assert isinstance(content, bytes) and len(content) > 0
    # A real .xlsx (zip) starts with the PK signature.
    assert content[:2] == b"PK"


# --------------------------------------------------------------------------- #
# 2. Summary tie-out: total_outstanding/overdue equal the API metrics EXACTLY.
# --------------------------------------------------------------------------- #
def test_summary_totals_equal_api_metrics_exactly(small_workbook, small_metrics):
    ws = small_workbook["Summary"]
    assert _summary_value(ws, "Total Outstanding") == small_metrics.total_outstanding
    assert _summary_value(ws, "Total Overdue") == small_metrics.total_overdue
    # Header block names the source file + as_of.
    assert _summary_value(ws, "Source File") == "ar_small.xlsx"
    assert _summary_value(ws, "As Of") == AS_OF.isoformat()
    # KPI + overall-bucket rows present.
    assert _summary_value(ws, "Customer Count") == small_metrics.customer_count
    assert _summary_value(ws, "Row Count") == small_metrics.row_count
    assert _summary_value(ws, "Worst Bucket") == small_metrics.worst_bucket
    assert _summary_value(ws, "Current") == small_metrics.bucket_totals.current
    assert _summary_value(ws, "90+") == small_metrics.bucket_totals.b_90_plus


def test_summary_money_cells_use_indian_currency_format(small_workbook):
    ws = small_workbook["Summary"]
    for row in ws.iter_rows():
        if row[0].value == "Total Outstanding":
            assert "₹" in row[1].number_format
            assert "#,##,##,##0.00" in row[1].number_format
            return
    raise AssertionError("Total Outstanding row not found")


# --------------------------------------------------------------------------- #
# 3. Customers sheet: EVERY customer, ranked by overdue desc, Beacon first.
# --------------------------------------------------------------------------- #
def test_customers_sheet_ranked_by_overdue_desc(small_workbook, small_metrics):
    ws = small_workbook["Customers"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Customer"
    assert "Overdue" in headers
    assert "Weighted Avg DPD" in headers

    data = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value is not None]
    # Uncapped: every customer present.
    assert len(data) == len(small_metrics.customer_breakdown) == 5
    # Ranked by overdue desc -> the deep-90+ customer is first.
    assert data[0] == "Beacon & Co"

    # Verify the overdue column is genuinely non-increasing down the sheet.
    overdue_col = headers.index("Overdue") + 1
    overdue_values = [
        ws.cell(row=r, column=overdue_col).value
        for r in range(2, 2 + len(data))
    ]
    assert overdue_values == sorted(overdue_values, reverse=True)


def test_customers_null_weighted_avg_renders_as_dash(small_workbook):
    ws = small_workbook["Customers"]
    headers = [c.value for c in ws[1]]
    cust_col = headers.index("Customer") + 1
    wavg_col = headers.index("Weighted Avg DPD") + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=cust_col).value == "Zenith Ltd":
            # Zenith is all-current -> weighted-avg is undefined -> em dash.
            assert ws.cell(row=r, column=wavg_col).value == "—"
            return
    raise AssertionError("Zenith Ltd not found in Customers sheet")


# --------------------------------------------------------------------------- #
# 4. Employees sheet: HoD Name column + every employee.
# --------------------------------------------------------------------------- #
def test_employees_sheet_has_hod_name_column(small_workbook, small_metrics):
    ws = small_workbook["Employees"]
    headers = [c.value for c in ws[1]]
    assert "HoD Name" in headers
    assert headers[0] == "Employee"

    data = [row[0].value for row in ws.iter_rows(min_row=2) if row[0].value is not None]
    assert len(data) == len(small_metrics.employees) == 3
    # Ranked desc by outstanding, blank-employee row present.
    assert data == ["Ravi", "Priya", "(blank)"]


# --------------------------------------------------------------------------- #
# 5. Flagged Rows sheet: EVERY flagged row + summary-excluded note.
# --------------------------------------------------------------------------- #
def test_flagged_rows_sheet_lists_all_flags_and_excluded_note(small_workbook, small_metrics):
    ws = small_workbook["Flagged Rows"]
    # Data rows = rows whose first cell is an integer row_index.
    data_rows = [
        row for row in ws.iter_rows(min_row=1, values_only=True)
        if isinstance(row[0], int)
    ]
    assert len(data_rows) == len(small_metrics.data_quality.rows)

    # The summary-rows-excluded count is shown (0 for ar_small — no embedded totals).
    excluded = small_metrics.data_quality.by_reason.get("summary_row_excluded", 0)
    note_text = " ".join(
        str(c.value) for c in (cell for r in ws.iter_rows() for cell in r) if c.value
    )
    assert "summary/total rows excluded" in note_text
    assert f"{excluded} summary/total rows excluded" in note_text
